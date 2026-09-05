"""The import template, and the contract it publishes."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.excel_template import (
    COLUMN_NAMES,
    COLUMNS,
    INSTRUCTIONS_SHEET,
    SHEET_NAME,
    TEMPLATE_FILENAME,
    TEMPLATE_VERSION,
    VERSION_LABEL_CELL,
    VERSION_VALUE_CELL,
    XLSX_MEDIA_TYPE,
    build_template_workbook,
)

# The playbook's eleven columns, written out longhand on purpose. If someone
# renames or reorders a column, this is what fails -- which is the entire point
# of calling the template a public contract. Copying COLUMN_NAMES here would
# make the test agree with the code by construction and assert nothing.
PLAYBOOK_COLUMNS = [
    "invoice_number",
    "vendor",
    "invoice_date",
    "item",
    "quantity",
    "unit_price",
    "tax_rate",
    "currency",
    "declared_subtotal",
    "declared_tax",
    "declared_total",
]


@pytest.fixture
def workbook():
    return load_workbook(BytesIO(build_template_workbook()))


# ------------------------------------------------------------ the contract


def test_columns_match_the_published_contract():
    assert list(COLUMN_NAMES) == PLAYBOOK_COLUMNS


def test_header_row_matches_the_contract(workbook):
    header = [cell.value for cell in workbook[SHEET_NAME][1]]
    assert header == PLAYBOOK_COLUMNS


# -------------------------------------------------------------- the workbook


def test_bytes_open_as_a_real_workbook(workbook):
    assert workbook.sheetnames == [SHEET_NAME, INSTRUCTIONS_SHEET]


def test_data_sheet_has_headers_and_nothing_else(workbook):
    """Anything left in the data sheet is something a user can import by mistake."""
    assert workbook[SHEET_NAME].max_row == 1


def test_version_is_readable_from_the_documented_cell(workbook):
    instructions = workbook[INSTRUCTIONS_SHEET]
    assert instructions[VERSION_LABEL_CELL].value == "template_version"
    assert instructions[VERSION_VALUE_CELL].value == TEMPLATE_VERSION


def test_every_column_is_documented(workbook):
    """The Instructions sheet and the column list cannot drift apart."""
    text = {
        cell.value
        for row in workbook[INSTRUCTIONS_SHEET].iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    }
    for name in COLUMN_NAMES:
        assert name in text


def test_tax_rate_percentage_is_spelled_out(workbook):
    """The ambiguity Phase 15 had to resolve should not reach a user."""
    tax_rate = next(c for c in COLUMNS if c.name == "tax_rate")
    assert "19 means 19%" in tax_rate.meaning


def test_the_worked_example_shows_one_invoice_over_two_rows(workbook):
    """A repeated invoice_number is the thing users most often get wrong.

    Located relative to the "Example:" heading rather than by counting the
    invoice number across the sheet -- it also appears in the column reference
    as that column's example, so a naive count finds three.
    """
    instructions = workbook[INSTRUCTIONS_SHEET]
    heading = next(
        row
        for row in range(1, instructions.max_row + 1)
        if str(instructions.cell(row=row, column=1).value or "").startswith("Example:")
    )

    # heading, then a header row, then the two example rows
    first, second = heading + 2, heading + 3
    assert instructions.cell(row=first, column=1).value == "INV-2026-1001"
    assert instructions.cell(row=second, column=1).value == "INV-2026-1001"
    # same invoice, different lines
    assert (
        instructions.cell(row=first, column=4).value
        != instructions.cell(row=second, column=4).value
    )


# --------------------------------------------------------------- the endpoint


def test_download_returns_a_spreadsheet(client):
    response = client.get("/templates/invoice-import")

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE


def test_download_is_named_with_its_version(client):
    disposition = client.get("/templates/invoice-import").headers["content-disposition"]

    assert TEMPLATE_FILENAME in disposition
    assert f"v{TEMPLATE_VERSION}" in disposition
    assert disposition.startswith("attachment;")


def test_downloaded_bytes_are_a_usable_workbook(client):
    """End to end: what the endpoint returns is what openpyxl can open."""
    response = client.get("/templates/invoice-import")
    downloaded = load_workbook(BytesIO(response.content))

    assert [c.value for c in downloaded[SHEET_NAME][1]] == PLAYBOOK_COLUMNS


def test_docs_describe_a_spreadsheet_not_json(client):
    schema = client.get("/openapi.json").json()
    content = schema["paths"]["/templates/invoice-import"]["get"]["responses"]["200"][
        "content"
    ]

    assert XLSX_MEDIA_TYPE in content
    assert "application/json" not in content
