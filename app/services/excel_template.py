"""The invoice import template, and the contract it publishes.

The playbook calls this template "a documented public contract", and that is
taken literally here: the column list below is the single definition of what an
import file looks like. This module writes it into a workbook, and from Phase 17
the upload validator reads the same constants to check what arrives. There is no
second copy to fall out of step with.

That is also why the template is generated rather than committed as a binary. A
checked-in .xlsx would be a second source of truth that nothing keeps honest.
"""
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

SHEET_NAME = "Invoices"
INSTRUCTIONS_SHEET = "Instructions"

# Bump when the columns change. Phase 17 reads it back off an upload to tell a
# user their template is out of date, rather than failing on a missing column
# and leaving them to guess why.
TEMPLATE_VERSION = 1

# The version's location is part of the contract, so it is named here rather
# than hard-coded at both ends.
VERSION_LABEL_CELL = "A1"
VERSION_VALUE_CELL = "B1"

TEMPLATE_FILENAME = f"invoiceflow-invoice-import-template-v{TEMPLATE_VERSION}.xlsx"

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


@dataclass(frozen=True)
class TemplateColumn:
    name: str
    meaning: str
    format: str
    example: str


# The playbook's eleven columns, in its order. Two of these names are mappings
# rather than passthroughs, and the Instructions sheet says so:
#
#   item          is what the API calls description
#   declared_*    is what the API calls subtotal / tax / total. "Declared"
#                 because on a row they sit beside per-line values and are only
#                 what the sender claims -- Phase 20 checks them against the
#                 lines, using the reconciliation Phase 15 built.
COLUMNS: tuple[TemplateColumn, ...] = (
    TemplateColumn(
        "invoice_number",
        "Identifies the invoice. Together with vendor it groups rows.",
        "Text",
        "INV-2026-1001",
    ),
    TemplateColumn(
        "vendor",
        "Who issued the invoice. Part of the grouping key.",
        "Text",
        "ABC GmbH",
    ),
    TemplateColumn(
        "invoice_date",
        "Date on the invoice. May not be in the future.",
        "YYYY-MM-DD",
        "2026-01-15",
    ),
    TemplateColumn(
        "item",
        "What this line is for. One row per line item.",
        "Text",
        "Consulting, March",
    ),
    TemplateColumn(
        "quantity",
        "How many. Must be greater than 0. May be fractional.",
        "Number, up to 3 decimals",
        "2.5",
    ),
    TemplateColumn(
        "unit_price",
        "Price for one unit, before tax.",
        "Number, 2 decimals",
        "100.00",
    ),
    TemplateColumn(
        "tax_rate",
        "Tax percentage for this line. 19 means 19%, not 0.19.",
        "Number, 0 to 100",
        "19",
    ),
    TemplateColumn(
        "currency",
        "Currency of the invoice. EUR, USD or GBP.",
        "3-letter code, uppercase",
        "EUR",
    ),
    TemplateColumn(
        "declared_subtotal",
        "Invoice total before tax, as you state it. Checked against the lines.",
        "Number, 2 decimals",
        "250.00",
    ),
    TemplateColumn(
        "declared_tax",
        "Invoice tax, as you state it. Checked against the lines.",
        "Number, 2 decimals",
        "47.50",
    ),
    TemplateColumn(
        "declared_total",
        "Invoice total including tax, as you state it. Checked against the lines.",
        "Number, 2 decimals",
        "297.50",
    ),
)

COLUMN_NAMES: tuple[str, ...] = tuple(column.name for column in COLUMNS)

# A worked example for the Instructions sheet: one invoice, two lines. The
# invoice-level columns repeat on every row of the same invoice, which looks
# like a mistake until you have seen it done deliberately.
_EXAMPLE_ROWS = (
    (
        "INV-2026-1001", "ABC GmbH", "2026-01-15", "Consulting, March",
        "2", "100.00", "19", "EUR", "250.00", "47.50", "297.50",
    ),
    (
        "INV-2026-1001", "ABC GmbH", "2026-01-15", "Travel",
        "1", "50.00", "19", "EUR", "250.00", "47.50", "297.50",
    ),
)


def build_template_workbook() -> bytes:
    """Build the import template as .xlsx bytes.

    The Invoices sheet gets the header row and nothing else. Anything left in
    a data sheet is something a user can forget to delete and then import as a
    real invoice, so the worked example lives on the Instructions sheet where
    it cannot be mistaken for data.
    """
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(list(COLUMN_NAMES))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    _write_instructions(workbook.create_sheet(INSTRUCTIONS_SHEET))

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_instructions(sheet) -> None:
    sheet[VERSION_LABEL_CELL] = "template_version"
    sheet[VERSION_VALUE_CELL] = TEMPLATE_VERSION
    sheet[VERSION_LABEL_CELL].font = Font(bold=True)

    sheet["A3"] = "InvoiceFlow invoice import"
    sheet["A3"].font = Font(bold=True, size=14)
    sheet["A4"] = (
        "Fill in the 'Invoices' sheet. Every column is required on every row. "
        "One row per line item: an invoice with three lines takes three rows, "
        "repeating the invoice-level columns on each."
    )
    sheet["A4"].alignment = Alignment(wrap_text=True)

    header_row = 6
    for index, title in enumerate(("column", "meaning", "format", "example"), start=1):
        cell = sheet.cell(row=header_row, column=index, value=title)
        cell.font = Font(bold=True)

    for offset, column in enumerate(COLUMNS, start=1):
        row = header_row + offset
        sheet.cell(row=row, column=1, value=column.name)
        sheet.cell(row=row, column=2, value=column.meaning)
        sheet.cell(row=row, column=3, value=column.format)
        sheet.cell(row=row, column=4, value=column.example)

    example_header = header_row + len(COLUMNS) + 2
    sheet.cell(row=example_header, column=1, value="Example: one invoice, two lines")
    sheet.cell(row=example_header, column=1).font = Font(bold=True)

    for index, name in enumerate(COLUMN_NAMES, start=1):
        cell = sheet.cell(row=example_header + 1, column=index, value=name)
        cell.font = Font(bold=True)
    for offset, example_row in enumerate(_EXAMPLE_ROWS, start=2):
        for index, value in enumerate(example_row, start=1):
            sheet.cell(row=example_header + offset, column=index, value=value)

    for column_letter, width in (("A", 22), ("B", 58), ("C", 24), ("D", 20)):
        sheet.column_dimensions[column_letter].width = width
